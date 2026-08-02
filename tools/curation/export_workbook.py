"""Exports the citation work list as an Excel workbook a curator can open and
fill in directly, instead of a raw CSV.

    python export_workbook.py                  # every unsourced material value
    python export_workbook.py --material ldpe   # one material only (pilot use)

This is a front-end over the same data export_gaps.py produces -- it calls
export_gaps()/export_sources() from that module (via a temp directory) so the
data pulled from the database is identical, then lays it out as three sheets:

  Instructions -- what this file is, the workflow loop, and a table
                  explaining every column (opens first)
  Values       -- the GAPS_FIELDNAMES worksheet, with the pre-filled columns
                  greyed out and locked, and the curator columns validated
                  with dropdowns where the value set is closed
  Sources      -- the SOURCES_FIELDNAMES bibliography, pre-populated, with
                  dropdowns for kind/tier and room for new rows

curation/polypedia-curation.xlsx is gitignored, same as curation/*.csv -- see
README.md's "Pipeline" section. import_workbook.py is the other half: it
reads this workbook back into gaps.csv/sources.csv and hands off to the
existing, unmodified import_values.py.
"""
from __future__ import annotations

import argparse
import csv
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from common import (
    CSV_ENCODING,
    CURATION_DIR,
    GAPS_CURATOR_FIELDS,
    GAPS_EXPORT_FIELDS,
    GAPS_FIELDNAMES,
    SOURCES_FIELDNAMES,
    get_connection,
)
from export_gaps import export_gaps, export_sources

WORKBOOK_PATH = CURATION_DIR / "polypedia-curation.xlsx"

# Real enum values from the database (source_kind / source_tier types in
# db/migrations) -- queried directly rather than guessed, see README.
SOURCE_KIND_VALUES = [
    "handbook",
    "textbook",
    "standard",
    "datasheet",
    "journal_article",
    "encyclopedia",
    "website",
    "internal",
]
SOURCE_TIER_VALUES = [
    "peer_reviewed_handbook",
    "standard",
    "manufacturer_datasheet",
    "vendor_marketing",
    "community",
]
QUALIFIER_VALUES = ["<", ">", "~", ">=", "<="]

# Extra blank rows below the seeded sources so the curator has room to add
# new ones, with the dropdowns/formats already extended to cover them.
SOURCES_EXTRA_ROWS = 50

# ---------------------------------------------------------------------------
# Styling
# ---------------------------------------------------------------------------

PREFILLED_HEADER_FILL = PatternFill("solid", fgColor="808080")
PREFILLED_DATA_FILL = PatternFill("solid", fgColor="D9D9D9")
CURATOR_HEADER_FILL = PatternFill("solid", fgColor="1F6F54")
CURATOR_DATA_FILL = PatternFill("solid", fgColor="FFFFFF")
HEADER_FONT_LIGHT = Font(bold=True, color="FFFFFF")
HEADER_FONT_DARK = Font(bold=True, color="000000")
WRAP_TOP_LEFT = Alignment(wrap_text=True, vertical="top", horizontal="left")

# ---------------------------------------------------------------------------
# Per-column guidance, reusing docs/CURATION-GUIDE.md wording so the two
# documents never disagree.
# ---------------------------------------------------------------------------

# (column, filled_by, meaning, example)
COLUMN_GUIDANCE: list[tuple[str, str, str, str]] = [
    ("material_slug", "Pre-filled -- do not edit",
     "Identifies which material this row is about.", "ldpe"),
    ("property_key", "Pre-filled -- do not edit",
     "Identifies which property this row is about.", "density"),
    ("property_name_en", "Pre-filled (context)",
     "English property name, so you know what the row is.", "Density"),
    ("property_name_fa", "Pre-filled (context)",
     "Persian property name, so you know what the row is.", "چگالی"),
    ("unit", "Pre-filled (context)",
     "The unit your number must be in.", "g/cm³"),
    ("plausible_min", "Pre-filled (context)",
     "Sanity floor -- if your number falls outside plausible_min/plausible_max, "
     "something is wrong (often a unit mix-up).", "0.8"),
    ("plausible_max", "Pre-filled (context)",
     "Sanity ceiling -- see plausible_min.", "2.3"),
    ("current_value", "Pre-filled (context)",
     "What the old prototype claimed, for comparison. Not read on import.", "0.91 - 0.94"),
    ("value_min", "You fill in", "For a range: the bottom.", "0.910"),
    ("value_max", "You fill in", "For a range: the top.", "0.925"),
    ("value_typical", "You fill in",
     "For a single number, not a range. Use this OR value_min/value_max, not both.", "-110"),
    ("qualifier", "You fill in",
     "Only if the source says something like \"less than 0.01\" -> put \"<\" here "
     "and 0.01 in value_max.", "<"),
    ("source_key", "You fill in",
     "The short name from the Sources sheet identifying the book/document you read.",
     "brydson-plastics-materials"),
    ("page", "You fill in",
     "The page you actually read it on. Required unless table/figure/section is filled in.", "45"),
    ("table", "You fill in", "Instead of, or as well as, a page.", "3.2"),
    ("figure", "You fill in", "Instead of, or as well as, a page.", "5"),
    ("section", "You fill in", "Instead of, or as well as, a page.", "2.1.3"),
    ("test_method", "You fill in", "If the source states one, e.g. a standard code.", "ASTM D1238"),
    ("conditions", "You fill in", "If the source states test conditions.", "190C/2.16kg"),
    ("note_en", "You fill in", "Anything worth remembering, in English.",
     "Value from Table 3.2, LDPE film grade"),
    ("note_fa", "You fill in", "Anything worth remembering, in Persian.", "یادداشت"),
    ("confidence", "You fill in", "0 to 1. Leave blank for the default (0.9).", "0.9"),
    ("skip", "You fill in", "Put \"y\" to ignore this row for now.", "y"),
]

TIER_GUIDANCE = [
    ("peer_reviewed_handbook", "Handbooks, textbooks, encyclopedias."),
    ("standard", "ASTM, ISO, and similar standards bodies."),
    ("manufacturer_datasheet", "A producer's technical datasheet (PDF)."),
    ("vendor_marketing", "A sales brochure or marketing material."),
    ("community", "Wikipedia and similar community-maintained sources."),
]

QUALIFIER_GUIDANCE = [
    ("<", "The real value is less than the number given."),
    (">", "The real value is greater than the number given."),
    ("~", "The number given is approximate."),
    (">=", "The real value is greater than or equal to the number given."),
    ("<=", "The real value is less than or equal to the number given."),
]

COLUMN_WIDTHS: dict[str, int] = {
    "material_slug": 14,
    "property_key": 18,
    "property_name_en": 22,
    "property_name_fa": 28,
    "unit": 10,
    "plausible_min": 12,
    "plausible_max": 12,
    "current_value": 16,
    "value_min": 12,
    "value_max": 12,
    "value_typical": 13,
    "qualifier": 10,
    "source_key": 28,
    "page": 8,
    "table": 8,
    "figure": 8,
    "section": 10,
    "test_method": 14,
    "conditions": 18,
    "note_en": 32,
    "note_fa": 32,
    "confidence": 11,
    "skip": 7,
}


# ---------------------------------------------------------------------------
# Instructions sheet
# ---------------------------------------------------------------------------


def build_instructions_sheet(wb: Workbook) -> Worksheet:
    ws = wb.active
    ws.title = "Instructions"
    ws.sheet_view.rightToLeft = False

    row = 1

    def write(text: str, *, bold: bool = False, size: int = 11) -> None:
        nonlocal row
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = Font(bold=bold, size=size)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    write("Polypedia Curation Workbook", bold=True, size=16)
    write("")
    write(
        "This file is the fill-in-the-blanks front end for the citation campaign: "
        "109 property values in the database are missing a source, and this workbook "
        "is how a domain expert adds one without writing SQL.",
    )
    write("")
    write("The 4-step loop", bold=True, size=13)
    write("1. export      -- python tools/curation/export_workbook.py --material ldpe")
    write("2. fill        -- open the Values sheet below and fill in what you know")
    write(
        "3. dry-run     -- python tools/curation/import_workbook.py --dry-run "
        "(checks your work, writes nothing)"
    )
    write(
        "4. import      -- python tools/curation/import_workbook.py "
        "(writes to the database, all or nothing)"
    )
    write("")
    write(
        "THE RULE THAT MATTERS: never invent a page number. If you cannot find where "
        "a value comes from, leave the row blank. 'unsourced' is honest -- a made-up "
        "citation cannot be detected later and destroys the one thing that makes this "
        "project worth building.",
        bold=True,
    )
    write("")
    write(
        "You do not have to fill the whole file in one sitting. Blank rows are simply "
        "ignored -- do ten rows, import them, come back tomorrow."
    )
    write("")

    # --- Values column guidance table ---------------------------------
    write("Values sheet -- column by column", bold=True, size=13)
    header_row = row
    headers = ["Column", "Filled by", "What it means", "Example"]
    for col, text in enumerate(headers, start=1):
        c = ws.cell(row=header_row, column=col, value=text)
        c.font = HEADER_FONT_LIGHT
        c.fill = PREFILLED_HEADER_FILL
        c.alignment = WRAP_TOP_LEFT
    row += 1

    for column, filled_by, meaning, example in COLUMN_GUIDANCE:
        is_curator = filled_by.startswith("You fill in")
        values = [column, filled_by, meaning, example]
        for col, text in enumerate(values, start=1):
            c = ws.cell(row=row, column=col, value=text)
            c.alignment = WRAP_TOP_LEFT
            c.fill = CURATOR_DATA_FILL if is_curator else PREFILLED_DATA_FILL
        row += 1

    row += 1

    # --- tier guidance ---------------------------------------------------
    write("The 'tier' column (Sources sheet) -- how much to trust a source", bold=True, size=13)
    tier_header_row = row
    for col, text in enumerate(["tier value", "Meaning"], start=1):
        c = ws.cell(row=tier_header_row, column=col, value=text)
        c.font = HEADER_FONT_LIGHT
        c.fill = PREFILLED_HEADER_FILL
        c.alignment = WRAP_TOP_LEFT
    row += 1
    for tier, meaning in TIER_GUIDANCE:
        ws.cell(row=row, column=1, value=tier).alignment = WRAP_TOP_LEFT
        ws.cell(row=row, column=2, value=meaning).alignment = WRAP_TOP_LEFT
        row += 1

    row += 1

    # --- qualifier guidance -----------------------------------------------
    write("The 'qualifier' column (Values sheet) -- inequality symbols", bold=True, size=13)
    qual_header_row = row
    for col, text in enumerate(["qualifier", "Meaning"], start=1):
        c = ws.cell(row=qual_header_row, column=col, value=text)
        c.font = HEADER_FONT_LIGHT
        c.fill = PREFILLED_HEADER_FILL
        c.alignment = WRAP_TOP_LEFT
    row += 1
    for symbol, meaning in QUALIFIER_GUIDANCE:
        ws.cell(row=row, column=1, value=symbol).alignment = WRAP_TOP_LEFT
        ws.cell(row=row, column=2, value=meaning).alignment = WRAP_TOP_LEFT
        row += 1

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 55
    ws.column_dimensions["D"].width = 30
    ws.freeze_panes = "A2"
    return ws


# ---------------------------------------------------------------------------
# Values sheet
# ---------------------------------------------------------------------------


def _guidance_lookup() -> dict[str, tuple[str, str, str]]:
    return {col: (filled_by, meaning, example) for col, filled_by, meaning, example in COLUMN_GUIDANCE}


def build_values_sheet(wb: Workbook, gaps_rows: list[dict], sources_row_count: int) -> Worksheet:
    ws = wb.create_sheet("Values")
    ws.sheet_view.rightToLeft = False
    guidance = _guidance_lookup()

    for col_idx, field in enumerate(GAPS_FIELDNAMES, start=1):
        is_curator = field in GAPS_CURATOR_FIELDS
        cell = ws.cell(row=1, column=col_idx, value=field)
        cell.font = HEADER_FONT_LIGHT if is_curator else HEADER_FONT_DARK
        cell.fill = CURATOR_HEADER_FILL if is_curator else PREFILLED_HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
        cell.protection = Protection(locked=True)
        if is_curator:
            _, meaning, example = guidance[field]
            cell.comment = Comment(
                f"You fill this in. {meaning}\nExample: {example}\n"
                "See the Instructions sheet for the full column-by-column guide.",
                "Polypedia curation workbook",
            )
        width = COLUMN_WIDTHS.get(field, 14)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    numeric_fields = {"plausible_min", "plausible_max", "value_min", "value_max", "value_typical"}

    for row_offset, row in enumerate(gaps_rows, start=2):
        for col_idx, field in enumerate(GAPS_FIELDNAMES, start=1):
            is_curator = field in GAPS_CURATOR_FIELDS
            raw = row.get(field, "")
            value: object = raw
            if not is_curator and field in numeric_fields and raw not in (None, ""):
                try:
                    value = float(raw)
                except (TypeError, ValueError):
                    value = raw
            elif is_curator:
                # Curator columns must come back empty on a fresh export --
                # a pre-filled source_key/page here would be an invented
                # citation waiting to be imported.
                value = None
            elif raw == "":
                value = None

            cell = ws.cell(row=row_offset, column=col_idx, value=value)
            cell.fill = CURATOR_DATA_FILL if is_curator else PREFILLED_DATA_FILL
            cell.protection = Protection(locked=not is_curator)
            if field == "confidence":
                cell.number_format = "0.00"
            elif field in numeric_fields:
                cell.number_format = "General"

    last_data_row = len(gaps_rows) + 1

    # --- data validation dropdowns -----------------------------------
    if last_data_row >= 2:
        qualifier_col = get_column_letter(GAPS_FIELDNAMES.index("qualifier") + 1)
        skip_col = get_column_letter(GAPS_FIELDNAMES.index("skip") + 1)
        source_key_col = get_column_letter(GAPS_FIELDNAMES.index("source_key") + 1)

        dv_qualifier = DataValidation(
            type="list", formula1=f'"{",".join(QUALIFIER_VALUES)}"', allow_blank=True,
            showErrorMessage=True, errorTitle="Invalid qualifier",
            error="Use one of: " + ", ".join(QUALIFIER_VALUES) + " (or leave blank).",
        )
        ws.add_data_validation(dv_qualifier)
        dv_qualifier.add(f"{qualifier_col}2:{qualifier_col}{last_data_row}")

        dv_skip = DataValidation(
            type="list", formula1='"y"', allow_blank=True,
            showErrorMessage=True, errorTitle="Invalid skip value",
            error='Only "y" is accepted (or leave blank).',
        )
        ws.add_data_validation(dv_skip)
        dv_skip.add(f"{skip_col}2:{skip_col}{last_data_row}")

        # +1 header row, +SOURCES_EXTRA_ROWS room for sources the curator
        # adds later on the Sources sheet.
        source_last_row = 1 + sources_row_count + SOURCES_EXTRA_ROWS
        dv_source = DataValidation(
            type="list",
            formula1=f"Sources!$A$2:$A${source_last_row}",
            allow_blank=True,
            showErrorMessage=True, errorTitle="Unknown source_key",
            error="Pick a source_key from the Sources sheet, or add a new row there first.",
        )
        ws.add_data_validation(dv_source)
        dv_source.add(f"{source_key_col}2:{source_key_col}{last_data_row}")

    # Freeze below the header and after the identifying columns
    # (material_slug, property_key) so they stay visible while scrolling
    # right through the rest of the row.
    ws.freeze_panes = "C2"

    # Protect the pre-filled columns. No password is set -- this is a
    # guard-rail against an accidental edit, not a security boundary,
    # and an unlocked "Unprotect Sheet" keeps the file from becoming
    # awkward to work with if a curator ever needs to override it.
    ws.protection.sheet = True
    ws.protection.formatCells = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.sort = False
    ws.protection.autoFilter = False

    return ws


# ---------------------------------------------------------------------------
# Sources sheet
# ---------------------------------------------------------------------------


def build_sources_sheet(wb: Workbook, sources_rows: list[dict]) -> Worksheet:
    ws = wb.create_sheet("Sources")
    ws.sheet_view.rightToLeft = False

    widths = {
        "source_key": 28, "title": 40, "authors": 24, "publisher": 22,
        "edition": 12, "year": 8, "isbn": 16, "doi": 20, "url": 30,
        "kind": 16, "tier": 24,
    }
    for col_idx, field in enumerate(SOURCES_FIELDNAMES, start=1):
        cell = ws.cell(row=1, column=col_idx, value=field)
        cell.font = HEADER_FONT_LIGHT
        cell.fill = PREFILLED_HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
        ws.column_dimensions[get_column_letter(col_idx)].width = widths.get(field, 16)

    for row_offset, row in enumerate(sources_rows, start=2):
        for col_idx, field in enumerate(SOURCES_FIELDNAMES, start=1):
            raw = row.get(field, "")
            value: object = raw if raw != "" else None
            if field == "year" and value is not None:
                try:
                    value = int(value)
                except (TypeError, ValueError):
                    pass
            ws.cell(row=row_offset, column=col_idx, value=value)

    last_row = 1 + len(sources_rows) + SOURCES_EXTRA_ROWS
    kind_col = get_column_letter(SOURCES_FIELDNAMES.index("kind") + 1)
    tier_col = get_column_letter(SOURCES_FIELDNAMES.index("tier") + 1)

    dv_kind = DataValidation(
        type="list", formula1=f'"{",".join(SOURCE_KIND_VALUES)}"', allow_blank=True,
        showErrorMessage=True, errorTitle="Invalid kind",
        error="Use one of: " + ", ".join(SOURCE_KIND_VALUES),
    )
    ws.add_data_validation(dv_kind)
    dv_kind.add(f"{kind_col}2:{kind_col}{last_row}")

    dv_tier = DataValidation(
        type="list", formula1=f'"{",".join(SOURCE_TIER_VALUES)}"', allow_blank=True,
        showErrorMessage=True, errorTitle="Invalid tier",
        error="Use one of: " + ", ".join(SOURCE_TIER_VALUES),
    )
    ws.add_data_validation(dv_tier)
    dv_tier.add(f"{tier_col}2:{tier_col}{last_row}")

    ws.freeze_panes = "A2"
    return ws


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------


def _read_csv_rows(path: Path) -> list[dict]:
    with path.open(newline="", encoding=CSV_ENCODING) as f:
        return list(csv.DictReader(f))


def export_workbook(conn, material_slug: str | None, out_path: Path = WORKBOOK_PATH) -> tuple[int, int]:
    """Builds the workbook by reusing export_gaps()/export_sources() against
    a temp directory (so the on-disk data contract can never drift from the
    CSV pipeline), then lays the same rows out as three formatted sheets.
    """
    with tempfile.TemporaryDirectory() as tmp:
        tmp_gaps = Path(tmp) / "gaps.csv"
        tmp_sources = Path(tmp) / "sources.csv"
        export_gaps(conn, material_slug, tmp_gaps)
        export_sources(conn, tmp_sources)
        gaps_rows = _read_csv_rows(tmp_gaps)
        sources_rows = _read_csv_rows(tmp_sources)

    wb = Workbook()
    build_instructions_sheet(wb)
    build_values_sheet(wb, gaps_rows, len(sources_rows))
    build_sources_sheet(wb, sources_rows)
    wb.active = wb.sheetnames.index("Instructions")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return len(gaps_rows), len(sources_rows)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--material",
        metavar="SLUG",
        default=None,
        help="Export gaps for one material only, e.g. ldpe. Default: all materials.",
    )
    parser.add_argument("--out", type=Path, default=WORKBOOK_PATH, help="Output .xlsx path.")
    args = parser.parse_args()

    conn = get_connection()
    try:
        gap_count, source_count = export_workbook(conn, args.material, args.out)
        conn.rollback()  # read-only session; close the transaction cleanly
    finally:
        conn.close()

    scope = f"material={args.material}" if args.material else "all materials"
    print(f"Wrote {args.out} ({gap_count} value rows, {source_count} source rows, {scope})")


if __name__ == "__main__":
    main()
