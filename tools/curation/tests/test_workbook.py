"""Tests for the Excel front end (export_workbook.py / import_workbook.py)
over the CSV curation workflow.

Follows the conventions in test_integration.py / conftest.py: tests that
touch the database use the `db_conn` fixture (rolled back in teardown,
never committed) or run the CLI in a way that provably never commits
(--dry-run). Nothing here writes to the real curation/*.csv or
curation/polypedia-curation.xlsx -- every test uses tmp_path.

A pre-filled source_key/page in a fresh export is not a cosmetic bug -- it
is an invented citation waiting to be imported (see README.md and
test_curation.py's module docstring for why an earlier iteration of this
project shipped exactly that). test_curator_columns_are_empty_on_export
guards against a regression of that specific failure mode.
"""
from __future__ import annotations

import csv
import datetime
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

from common import (
    GAPS_CURATOR_FIELDS,
    GAPS_EXPORT_FIELDS,
    GAPS_FIELDNAMES,
    SOURCES_FIELDNAMES,
    get_connection,
)
from export_workbook import export_workbook
from import_workbook import WorkbookCellError, cell_text, workbook_to_csvs

TOOLS_CURATION = Path(__file__).resolve().parents[1]
PYTHON = sys.executable
TEST_SOURCE_TITLE = "__test_source__ workbook fixture"


# ---------------------------------------------------------------------------
# Export: shape
# ---------------------------------------------------------------------------


def test_export_produces_three_sheets_with_expected_headers(db_conn, tmp_path):
    out_path = tmp_path / "wb.xlsx"
    gap_count, source_count = export_workbook(db_conn, "ldpe", out_path)

    assert out_path.exists()
    wb = load_workbook(out_path)
    assert wb.sheetnames == ["Instructions", "Values", "Sources"]
    assert wb.active.title == "Instructions"

    values_ws = wb["Values"]
    header = [c.value for c in next(values_ws.iter_rows(min_row=1, max_row=1))]
    assert header == GAPS_FIELDNAMES

    sources_ws = wb["Sources"]
    sources_header = [c.value for c in next(sources_ws.iter_rows(min_row=1, max_row=1))]
    assert sources_header == SOURCES_FIELDNAMES

    # gap_count/source_count are whatever export_workbook found -- not a
    # fixed number, since real curation work changes both (citing a value
    # removes it from the gap count; adding a bibliography entry grows the
    # source count). Only their internal consistency with the sheet
    # contents is asserted here.
    assert gap_count == values_ws.max_row - 1
    assert source_count == sources_ws.max_row - 1


def test_export_all_materials_covers_both_seeded_materials(db_conn, tmp_path):
    out_path = tmp_path / "wb-all.xlsx"
    gap_count, _ = export_workbook(db_conn, None, out_path)

    wb = load_workbook(out_path)
    values_ws = wb["Values"]
    slugs = {
        row[0].value
        for row in values_ws.iter_rows(min_row=2)
        if row[0].value is not None
    }
    assert slugs == {"ldpe", "hdpe"}
    assert gap_count == values_ws.max_row - 1


# ---------------------------------------------------------------------------
# Export: curator columns must come back empty (the critical regression)
# ---------------------------------------------------------------------------


def test_curator_columns_are_empty_on_export(db_conn, tmp_path):
    out_path = tmp_path / "wb.xlsx"
    export_workbook(db_conn, "ldpe", out_path)

    wb = load_workbook(out_path)
    ws = wb["Values"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    curator_cols = [header.index(f) for f in GAPS_CURATOR_FIELDS]

    offending = []
    for row in ws.iter_rows(min_row=2):
        for idx in curator_cols:
            if row[idx].value not in (None, ""):
                offending.append((row[0].row, header[idx], row[idx].value))

    assert offending == [], (
        "curator columns must be blank on a fresh export -- a pre-filled "
        "source_key/page is an invented citation waiting to be imported"
    )

    # And the pre-filled/context columns should, conversely, actually be
    # filled in (not accidentally blanked by the same logic).
    export_cols = [header.index(f) for f in GAPS_EXPORT_FIELDS if f not in ("plausible_min", "plausible_max")]
    first_row = next(ws.iter_rows(min_row=2, max_row=2))
    for idx in export_cols:
        assert first_row[idx].value not in (None, ""), header[idx]


# ---------------------------------------------------------------------------
# Persian text
# ---------------------------------------------------------------------------


def test_persian_text_survives_export(db_conn, tmp_path):
    out_path = tmp_path / "wb.xlsx"
    export_workbook(db_conn, "ldpe", out_path)

    wb = load_workbook(out_path)
    ws = wb["Values"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    fa_col = header.index("property_name_fa")

    persian_values = [row[fa_col].value for row in ws.iter_rows(min_row=2) if row[fa_col].value]
    assert persian_values, "expected at least one property with Persian context text"
    for v in persian_values:
        assert "�" not in v
        assert any("؀" <= ch <= "ۿ" for ch in v)


def test_persian_note_round_trips_through_workbook_to_csv(db_conn, tmp_path):
    out_path = tmp_path / "wb.xlsx"
    export_workbook(db_conn, "ldpe", out_path)

    wb = load_workbook(out_path)
    ws = wb["Values"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    note_fa_col = header.index("note_fa") + 1  # 1-based for cell()
    persian_note = "چگالی بین دو مقدار است"
    ws.cell(row=2, column=note_fa_col, value=persian_note)
    wb.save(out_path)

    gaps_csv = tmp_path / "gaps.csv"
    sources_csv = tmp_path / "sources.csv"
    workbook_to_csvs(out_path, gaps_csv, sources_csv)

    with gaps_csv.open(encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    assert rows[0]["note_fa"] == persian_note


# ---------------------------------------------------------------------------
# Filled workbook -> CSV round trip
# ---------------------------------------------------------------------------


def test_filled_workbook_round_trips_to_csv(db_conn, tmp_path):
    out_path = tmp_path / "wb.xlsx"
    export_workbook(db_conn, "ldpe", out_path)

    wb = load_workbook(out_path)
    ws = wb["Values"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {name: header.index(name) + 1 for name in header}

    # Row 2: cite it against an existing seeded source.
    ws.cell(row=2, column=col["source_key"], value="polymer-handbook-4e")
    ws.cell(row=2, column=col["page"], value="45")
    ws.cell(row=2, column=col["value_typical"], value=0.92)
    ws.cell(row=2, column=col["confidence"], value=0.85)
    ws.cell(row=2, column=col["note_en"], value="from fixture")

    # Add a brand-new source on the Sources sheet.
    sources_ws = wb["Sources"]
    sheader = [c.value for c in next(sources_ws.iter_rows(min_row=1, max_row=1))]
    scol = {name: sheader.index(name) + 1 for name in sheader}
    new_row = sources_ws.max_row + 1
    sources_ws.cell(row=new_row, column=scol["source_key"], value="test-workbook-source")
    sources_ws.cell(row=new_row, column=scol["title"], value=TEST_SOURCE_TITLE)
    sources_ws.cell(row=new_row, column=scol["kind"], value="handbook")
    sources_ws.cell(row=new_row, column=scol["tier"], value="community")

    wb.save(out_path)

    gaps_csv = tmp_path / "gaps.csv"
    sources_csv = tmp_path / "sources.csv"
    gap_count, source_count = workbook_to_csvs(out_path, gaps_csv, sources_csv)

    with gaps_csv.open(encoding="utf-8-sig", newline="") as f:
        gaps_rows = list(csv.DictReader(f))
    with sources_csv.open(encoding="utf-8-sig", newline="") as f:
        sources_rows = list(csv.DictReader(f))

    assert gap_count == len(gaps_rows)
    assert source_count == len(sources_rows)

    first = gaps_rows[0]
    assert first["source_key"] == "polymer-handbook-4e"
    assert first["page"] == "45"
    assert first["value_typical"] == "0.92"
    assert first["confidence"] == "0.85"
    assert first["note_en"] == "from fixture"
    # untouched curator columns should still come back blank, not "None"
    assert first["qualifier"] == ""
    assert first["skip"] == ""

    added = [r for r in sources_rows if r["source_key"] == "test-workbook-source"]
    assert len(added) == 1
    assert added[0]["title"] == TEST_SOURCE_TITLE
    assert added[0]["kind"] == "handbook"
    assert added[0]["tier"] == "community"


# ---------------------------------------------------------------------------
# Date-typed cell -> readable error, not a crash
# ---------------------------------------------------------------------------


def test_date_typed_cell_is_rejected_with_readable_error_not_a_crash():
    """Simulates the classic Excel quirk: a curator types a bare number
    into a locator column and Excel silently reinterprets it as a date.
    openpyxl hands that back as a datetime -- cell_text() must turn it into
    a plain-language, cell-addressed error rather than raising a bare
    TypeError/ValueError or writing the wrong thing to the CSV.
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Values"
    for idx, field in enumerate(GAPS_FIELDNAMES, start=1):
        ws.cell(row=1, column=idx, value=field)
    page_col = GAPS_FIELDNAMES.index("page") + 1
    ws.cell(row=2, column=1, value="ldpe")
    ws.cell(row=2, column=2, value="density")
    ws.cell(row=2, column=page_col, value=datetime.date(2026, 4, 1))

    with pytest.raises(WorkbookCellError) as exc:
        cell_text("Values", ws.cell(row=2, column=page_col))
    message = str(exc.value)
    assert "Values!N2" in message
    assert "looks like a date" in message
    assert "2026-04-01" in message
    assert "Format the column as Text or Number" in message


def test_date_typed_cell_via_workbook_to_csvs_reports_and_does_not_crash(db_conn, tmp_path):
    out_path = tmp_path / "wb.xlsx"
    export_workbook(db_conn, "ldpe", out_path)

    wb = load_workbook(out_path)
    ws = wb["Values"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    page_col = header.index("page") + 1
    ws.cell(row=2, column=page_col, value=datetime.datetime(2026, 4, 1))
    wb.save(out_path)

    gaps_csv = tmp_path / "gaps.csv"
    sources_csv = tmp_path / "sources.csv"

    with pytest.raises(SystemExit) as exc:
        workbook_to_csvs(out_path, gaps_csv, sources_csv)
    message = str(exc.value)
    assert "looks like a date" in message
    assert "2026-04-01" in message
    assert "Values!" in message
    # Nothing should have been written -- reject rather than guess.
    assert not gaps_csv.exists()
    assert not sources_csv.exists()


# ---------------------------------------------------------------------------
# None vs empty string, and stray whitespace on a text-formatted number
# ---------------------------------------------------------------------------


def test_blank_cell_and_whitespace_padded_text_cell_normalize_cleanly():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value=None)
    ws.cell(row=1, column=2, value="  0.925  ")
    ws.cell(row=1, column=3, value="")

    assert cell_text("Values", ws.cell(row=1, column=1)) == ""
    assert cell_text("Values", ws.cell(row=1, column=2)) == "0.925"
    assert cell_text("Values", ws.cell(row=1, column=3)) == ""


def test_numeric_cell_renders_without_spurious_trailing_zero():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value=300.0)
    ws.cell(row=1, column=2, value=0.925)

    assert cell_text("Values", ws.cell(row=1, column=1)) == "300"
    assert cell_text("Values", ws.cell(row=1, column=2)) == "0.925"


# ---------------------------------------------------------------------------
# End-to-end dry-run through import_workbook.py's own CLI (never commits)
# ---------------------------------------------------------------------------


def test_import_workbook_dry_run_end_to_end(tmp_path):
    """Exercises the real import_workbook.py CLI: export a workbook, fill
    one row against a seeded source, run --dry-run, and confirm it reports
    a row ready to import without writing anything to the database.
    """
    conn = get_connection()
    try:
        wb_path = tmp_path / "wb.xlsx"
        from export_workbook import export_workbook as _export

        _export(conn, "ldpe", wb_path)
    finally:
        conn.rollback()
        conn.close()

    wb = load_workbook(wb_path)
    ws = wb["Values"]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col = {name: header.index(name) + 1 for name in header}
    ws.cell(row=2, column=col["source_key"], value="polymer-handbook-4e")
    ws.cell(row=2, column=col["page"], value="7")
    wb.save(wb_path)

    material_slug = ws.cell(row=2, column=col["material_slug"]).value
    property_key = ws.cell(row=2, column=col["property_key"]).value

    gaps_csv = tmp_path / "gaps.csv"
    sources_csv = tmp_path / "sources.csv"

    result = subprocess.run(
        [
            PYTHON, str(TOOLS_CURATION / "import_workbook.py"),
            "--workbook", str(wb_path),
            "--gaps-csv", str(gaps_csv),
            "--sources-csv", str(sources_csv),
            "--dry-run",
        ],
        cwd=TOOLS_CURATION, capture_output=True, text=True, timeout=30,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert "DRY RUN" in result.stdout
    assert "1 row(s)" in result.stdout or "1 to import" in result.stdout

    # Confirm the database is untouched.
    check_conn = get_connection()
    try:
        with check_conn.cursor() as cur:
            cur.execute(
                """
                SELECT pv.status FROM property_value pv
                JOIN material m ON m.id = pv.subject_id AND pv.subject_type = 'material'
                JOIN property_definition pd ON pd.id = pv.property_id
                WHERE m.slug = %s AND pd.key = %s;
                """,
                (material_slug, property_key),
            )
            assert cur.fetchone()[0] == "unsourced"
    finally:
        check_conn.rollback()
        check_conn.close()
